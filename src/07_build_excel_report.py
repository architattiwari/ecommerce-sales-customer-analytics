"""Step 7: build reports/ecommerce_analysis.xlsx.

The workbook keeps a compact summary table (Dashboard_Data) and uses Excel formulas
(SUMIFS, COUNTIFS, INDEX/MATCH, IF, RANK) for everything else, so the numbers update
if you change the data.
Blue numbers are values copied from Python, black numbers are formulas.
"""
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from common import DELIVERY_SLA_DAYS, PROCESSED_DIR, REPORTS_DIR, build_sales_table, load_processed

FONT = "Arial"
HEADER_FILL = PatternFill("solid", start_color="1F4E78")
INPUT_FILL = PatternFill("solid", start_color="FFF2CC")
BLUE = Font(name=FONT, size=10, color="0000FF")
BLACK = Font(name=FONT, size=10)
BOLD = Font(name=FONT, size=10, bold=True)
TITLE = Font(name=FONT, size=14, bold=True)
NOTE = Font(name=FONT, size=9, italic=True, color="595959")

MONEY = "#,##0"
PERCENT = "0.0%"


def write_header(ws, row, headers, start_col=1):
    for i, text in enumerate(headers):
        cell = ws.cell(row=row, column=start_col + i, value=text)
        cell.font = Font(name=FONT, size=10, bold=True, color="FFFFFF")
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)


def put(ws, row, col, value, number_format=None, font=None):
    cell = ws.cell(row=row, column=col, value=value)
    is_formula = isinstance(value, str) and value.startswith("=")
    cell.font = font or (BLACK if (is_formula or isinstance(value, str)) else BLUE)
    if number_format:
        cell.number_format = number_format
    return cell


def set_widths(ws, widths):
    for i, width in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = width


# ---------------------------------------------------------------------------
# Sheets
# ---------------------------------------------------------------------------
def build_dashboard_data(wb, sales):
    ws = wb.create_sheet("Dashboard_Data")
    grouped = sales.groupby(["order_month", "category", "shipping_state"]).agg(
        order_lines=("order_id", "size"), quantity=("quantity", "sum"), gross_sales=("gross_sales", "sum"),
        discount_amount=("discount_amount", "sum"), revenue=("revenue", "sum"), cost=("cost", "sum"),
        profit=("profit", "sum")).reset_index()
    write_header(ws, 1, ["Month", "Category", "State", "Order_Lines", "Quantity", "Gross_Sales",
                         "Discount_Amount", "Revenue", "Cost", "Profit"])
    for r, row in enumerate(grouped.itertuples(index=False), start=2):
        put(ws, r, 1, row.order_month.to_pydatetime(), "mmm yyyy", BLUE)
        put(ws, r, 2, row.category, font=BLUE)
        put(ws, r, 3, row.shipping_state, font=BLUE)
        for c, value in enumerate(row[3:], start=4):
            put(ws, r, c, round(float(value), 2), MONEY)
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:J{len(grouped) + 1}"
    set_widths(ws, [12, 24, 20, 12, 10, 14, 16, 14, 14, 14])
    return len(grouped) + 1


def build_monthly_sales(wb, sales, last_data_row):
    ws = wb.create_sheet("Monthly_Sales")
    months = sorted(sales["order_month"].unique())
    orders_per_month = sales.groupby("order_month")["order_id"].nunique()
    write_header(ws, 1, ["Month", "Orders (from Python)", "Revenue", "Profit", "Profit Margin", "Avg Order Value",
                         "Growth vs Last Month", "Cumulative Revenue"])
    dd = "Dashboard_Data"
    for i, month in enumerate(months):
        r = i + 2
        put(ws, r, 1, pd.Timestamp(month).to_pydatetime(), "mmm yyyy", BLUE)
        put(ws, r, 2, int(orders_per_month[month]), MONEY)
        put(ws, r, 3, f"=SUMIFS({dd}!$H$2:$H${last_data_row},{dd}!$A$2:$A${last_data_row},A{r})", MONEY)
        put(ws, r, 4, f"=SUMIFS({dd}!$J$2:$J${last_data_row},{dd}!$A$2:$A${last_data_row},A{r})", MONEY)
        put(ws, r, 5, f"=IF(C{r}=0,0,D{r}/C{r})", PERCENT)
        put(ws, r, 6, f"=IF(B{r}=0,0,C{r}/B{r})", MONEY)
        put(ws, r, 7, "" if i == 0 else f"=IF(C{r - 1}=0,0,C{r}/C{r - 1}-1)", PERCENT)
        put(ws, r, 8, f"=SUM($C$2:C{r})", MONEY)
    total_row = len(months) + 2
    put(ws, total_row, 1, "Total", font=BOLD)
    put(ws, total_row, 2, f"=SUM(B2:B{total_row - 1})", MONEY, BOLD)
    put(ws, total_row, 3, f"=SUM(C2:C{total_row - 1})", MONEY, BOLD)
    put(ws, total_row, 4, f"=SUM(D2:D{total_row - 1})", MONEY, BOLD)
    put(ws, total_row, 5, f"=IF(C{total_row}=0,0,D{total_row}/C{total_row})", PERCENT, BOLD)
    put(ws, total_row, 6, f"=IF(B{total_row}=0,0,C{total_row}/B{total_row})", MONEY, BOLD)
    put(ws, total_row + 2, 1, "Orders are distinct order counts, so they cannot be rebuilt from the line-level Dashboard_Data table.", font=NOTE)
    ws.freeze_panes = "A2"
    set_widths(ws, [12, 14, 15, 15, 14, 16, 18, 18])
    return total_row


def build_category_analysis(wb, sales, last_data_row):
    ws = wb.create_sheet("Category_Analysis")
    dd = "Dashboard_Data"
    categories = sales.groupby("category")["revenue"].sum().sort_values(ascending=False).index.tolist()
    first, last = 4, 4 + len(categories) - 1
    total = last + 1

    put(ws, 1, 1, "Category performance (valid sales only)", font=TITLE)
    write_header(ws, 3, ["Category", "Revenue", "Profit", "Units Sold", "Profit Margin", "Revenue Share",
                         "Revenue Rank", "Avg Discount", "Margin Check"])
    for i, category in enumerate(categories):
        r = first + i
        put(ws, r, 1, category, font=BLACK)
        put(ws, r, 2, f"=SUMIFS({dd}!$H$2:$H${last_data_row},{dd}!$B$2:$B${last_data_row},$A{r})", MONEY)
        put(ws, r, 3, f"=SUMIFS({dd}!$J$2:$J${last_data_row},{dd}!$B$2:$B${last_data_row},$A{r})", MONEY)
        put(ws, r, 4, f"=SUMIFS({dd}!$E$2:$E${last_data_row},{dd}!$B$2:$B${last_data_row},$A{r})", MONEY)
        put(ws, r, 5, f"=IF(B{r}=0,0,C{r}/B{r})", PERCENT)
        put(ws, r, 6, f"=B{r}/$B${total}", PERCENT)
        put(ws, r, 7, f"=RANK(B{r},$B${first}:$B${last})", "0")
        put(ws, r, 8, f"=SUMIFS({dd}!$G$2:$G${last_data_row},{dd}!$B$2:$B${last_data_row},$A{r})"
                      f"/SUMIFS({dd}!$F$2:$F${last_data_row},{dd}!$B$2:$B${last_data_row},$A{r})", PERCENT)
        put(ws, r, 9, f'=IF(E{r}<$E${total},"Below average margin","OK")')
    put(ws, total, 1, "Total", font=BOLD)
    put(ws, total, 2, f"=SUM(B{first}:B{last})", MONEY, BOLD)
    put(ws, total, 3, f"=SUM(C{first}:C{last})", MONEY, BOLD)
    put(ws, total, 4, f"=SUM(D{first}:D{last})", MONEY, BOLD)
    put(ws, total, 5, f"=IF(B{total}=0,0,C{total}/B{total})", PERCENT, BOLD)

    # pivot-style table: category x month
    months = sorted(sales["order_month"].unique())
    matrix_header = total + 3
    put(ws, matrix_header - 1, 1, "Revenue by category and month (pivot-style, SUMIFS)", font=BOLD)
    write_header(ws, matrix_header, ["Category"] + [""] * len(months) + ["Total"])
    for j, month in enumerate(months):
        cell = put(ws, matrix_header, 2 + j, pd.Timestamp(month).to_pydatetime(), "mmm", Font(name=FONT, size=10, bold=True, color="FFFFFF"))
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(horizontal="center")
    for i, category in enumerate(categories):
        r = matrix_header + 1 + i
        put(ws, r, 1, category, font=BLACK)
        for j in range(len(months)):
            col = get_column_letter(2 + j)
            put(ws, r, 2 + j, f"=SUMIFS({dd}!$H$2:$H${last_data_row},{dd}!$B$2:$B${last_data_row},$A{r},"
                              f"{dd}!$A$2:$A${last_data_row},{col}${matrix_header})", MONEY)
        put(ws, r, 2 + len(months), f"=SUM(B{r}:{get_column_letter(1 + len(months))}{r})", MONEY, BOLD)
    set_widths(ws, [24] + [13] * 13)
    return first, last, total


def build_product_analysis(wb, product_summary):
    ws = wb.create_sheet("Product_Analysis")
    header_row, first, last = 11, 12, 11 + len(product_summary)
    low_revenue_cutoff = float(product_summary.loc[product_summary["revenue"] > 0, "revenue"].quantile(0.10))

    put(ws, 1, 1, "Product performance (valid sales only)", font=TITLE)
    put(ws, 3, 1, "Product lookup", font=BOLD)
    put(ws, 4, 1, "Enter a product ID:")
    put(ws, 4, 2, "P0092", font=BLUE).fill = INPUT_FILL
    lookups = [("Product name", "B"), ("Category", "C"), ("Revenue", "G"), ("Profit", "H")]
    for i, (label, column) in enumerate(lookups):
        put(ws, 5 + i, 1, label)
        put(ws, 5 + i, 2, f'=IFERROR(INDEX(${column}${first}:${column}${last},MATCH($B$4,$A${first}:$A${last},0)),"Not found")',
            MONEY if label in ("Revenue", "Profit") else None)
    put(ws, 9, 1, f"Excel 365 alternative: =XLOOKUP(B4, A{first}:A{last}, B{first}:B{last})   (INDEX/MATCH is used above because it works in every Excel version)", font=NOTE)
    put(ws, 4, 4, "Low revenue cutoff (10th percentile of selling products, from Python):")
    put(ws, 4, 8, round(low_revenue_cutoff, 2), MONEY).fill = INPUT_FILL

    write_header(ws, header_row, ["Product ID", "Product Name", "Category", "Subcategory", "List Price", "Units Sold",
                                  "Revenue", "Profit", "Profit Margin", "Avg Selling Price", "Status"])
    for i, row in enumerate(product_summary.itertuples(index=False)):
        r = first + i
        put(ws, r, 1, row.product_id, font=BLACK)
        put(ws, r, 2, row.product_name, font=BLACK)
        put(ws, r, 3, row.category, font=BLACK)
        put(ws, r, 4, row.subcategory, font=BLACK)
        put(ws, r, 5, float(row.unit_price), MONEY)
        put(ws, r, 6, int(row.quantity_sold), MONEY)
        put(ws, r, 7, round(float(row.revenue), 2), MONEY)
        put(ws, r, 8, round(float(row.profit), 2), MONEY)
        put(ws, r, 9, f"=IF(G{r}=0,0,H{r}/G{r})", PERCENT)
        put(ws, r, 10, f"=IF(F{r}=0,0,G{r}/F{r})", MONEY)
        put(ws, r, 11, f'=IF(G{r}=0,"No sales",IF(H{r}<0,"Loss-making",IF(G{r}<$H$4,"Low revenue","OK")))')

    put(ws, 5, 4, "Products needing attention:", font=BOLD)
    for i, status in enumerate(["No sales", "Loss-making", "Low revenue"]):
        put(ws, 6 + i, 4, status)
        put(ws, 6 + i, 5, f"=COUNTIFS($K${first}:$K${last},D{6 + i})", MONEY)
    ws.freeze_panes = f"A{first}"
    ws.auto_filter.ref = f"A{header_row}:K{last}"
    set_widths(ws, [22, 34, 22, 22, 12, 11, 14, 14, 13, 16, 14])
    return first, last


def build_customer_analysis(wb, customer_summary):
    ws = wb.create_sheet("Customer_Analysis")
    header_row, first = 24, 25
    last = first + len(customer_summary) - 1

    def rng(column):
        return f"${column}${first}:${column}${last}"

    put(ws, 1, 1, "Customer analysis", font=TITLE)
    put(ws, 3, 1, "Value segments (COUNTIFS / SUMIFS on the customer table below)", font=BOLD)
    write_header(ws, 4, ["Value Segment", "Customers", "Revenue", "Revenue per Customer", "% of Customers", "% of Revenue", "Avg Orders"])
    segments = ["High Value", "Regular", "Occasional", "At Risk", "Inactive"]
    for i, segment in enumerate(segments):
        r = 5 + i
        put(ws, r, 1, segment, font=BLACK)
        put(ws, r, 2, f"=COUNTIFS({rng('D')},A{r})", MONEY)
        put(ws, r, 3, f"=SUMIFS({rng('F')},{rng('D')},A{r})", MONEY)
        put(ws, r, 4, f"=IF(B{r}=0,0,C{r}/B{r})", MONEY)
        put(ws, r, 5, f"=B{r}/$B$10", PERCENT)
        put(ws, r, 6, f"=C{r}/$C$10", PERCENT)
        put(ws, r, 7, f"=IF(B{r}=0,0,SUMIFS({rng('E')},{rng('D')},A{r})/B{r})", "0.0")
    put(ws, 10, 1, "Total", font=BOLD)
    put(ws, 10, 2, "=SUM(B5:B9)", MONEY, BOLD)
    put(ws, 10, 3, "=SUM(C5:C9)", MONEY, BOLD)

    put(ws, 12, 1, "Plan segments (chosen by the customer at signup)", font=BOLD)
    write_header(ws, 13, ["Plan Segment", "Customers", "Revenue", "Revenue per Customer"])
    for i, plan in enumerate(["Standard", "Premium", "Student", "Business"]):
        r = 14 + i
        put(ws, r, 1, plan, font=BLACK)
        put(ws, r, 2, f"=COUNTIFS({rng('C')},A{r})", MONEY)
        put(ws, r, 3, f"=SUMIFS({rng('F')},{rng('C')},A{r})", MONEY)
        put(ws, r, 4, f"=IF(B{r}=0,0,C{r}/B{r})", MONEY)

    put(ws, 19, 1, "Repeat customers (2+ valid orders)", font=BOLD)
    put(ws, 20, 1, "Customers with a sale")
    put(ws, 20, 2, f'=COUNTIFS({rng("E")},">=1")', MONEY)
    put(ws, 21, 1, "Repeat customers")
    put(ws, 21, 2, f'=COUNTIFS({rng("E")},">=2")', MONEY)
    put(ws, 22, 1, "Repeat customer rate")
    put(ws, 22, 2, "=IF(B20=0,0,B21/B20)", PERCENT)

    write_header(ws, header_row, ["Customer ID", "State", "Plan Segment", "Value Segment", "Orders", "Revenue",
                                  "Days Since Last Order", "CLV Estimate (3 yr)"])
    columns = ["customer_id", "state", "customer_segment", "value_segment", "total_orders", "total_revenue",
               "recency_days", "clv_estimate"]
    for i, row in enumerate(customer_summary[columns].itertuples(index=False)):
        r = first + i
        for c, value in enumerate(row, start=1):
            if pd.isna(value):
                continue
            if c <= 4:
                put(ws, r, c, value, font=BLACK)
            else:
                put(ws, r, c, round(float(value), 2), MONEY)
    ws.freeze_panes = f"A{first}"
    set_widths(ws, [30, 20, 16, 16, 16, 16, 20, 20])
    return first, last


def build_geography(wb, sales, last_data_row):
    ws = wb.create_sheet("Geography")
    dd = "Dashboard_Data"
    by_state = sales.groupby("shipping_state").agg(orders=("order_id", "nunique"), revenue=("revenue", "sum"))
    by_state = by_state.sort_values("revenue", ascending=False)
    first, last = 4, 3 + len(by_state)
    total = last + 1

    put(ws, 1, 1, "Revenue by shipping state (valid sales only)", font=TITLE)
    write_header(ws, 3, ["State", "Orders (from Python)", "Revenue", "Profit", "Profit Margin", "Avg Order Value",
                         "Revenue Share", "Cumulative Share", "Revenue Rank", "Market Size"])
    for i, (state, row) in enumerate(by_state.iterrows()):
        r = first + i
        put(ws, r, 1, state, font=BLACK)
        put(ws, r, 2, int(row["orders"]), MONEY)
        put(ws, r, 3, f"=SUMIFS({dd}!$H$2:$H${last_data_row},{dd}!$C$2:$C${last_data_row},$A{r})", MONEY)
        put(ws, r, 4, f"=SUMIFS({dd}!$J$2:$J${last_data_row},{dd}!$C$2:$C${last_data_row},$A{r})", MONEY)
        put(ws, r, 5, f"=IF(C{r}=0,0,D{r}/C{r})", PERCENT)
        put(ws, r, 6, f"=IF(B{r}=0,0,C{r}/B{r})", MONEY)
        put(ws, r, 7, f"=C{r}/$C${total}", PERCENT)
        put(ws, r, 8, f"=SUM($G${first}:G{r})", PERCENT)
        put(ws, r, 9, f"=RANK(C{r},$C${first}:$C${last})", "0")
        put(ws, r, 10, f'=IF(G{r}>=0.08,"Large",IF(G{r}>=0.03,"Medium","Small"))')
    put(ws, total, 1, "Total", font=BOLD)
    put(ws, total, 2, f"=SUM(B{first}:B{last})", MONEY, BOLD)
    put(ws, total, 3, f"=SUM(C{first}:C{last})", MONEY, BOLD)
    put(ws, total, 4, f"=SUM(D{first}:D{last})", MONEY, BOLD)
    ws.freeze_panes = "A4"
    set_widths(ws, [22, 14, 15, 15, 13, 16, 14, 16, 12, 13])
    return first, last, total


def build_kpi_summary(wb, orders, reviews, refs):
    ws = wb["KPI_Summary"]
    dd_last = refs["dd_last"]
    cust_first, cust_last = refs["cust_first"], refs["cust_last"]
    delivered = orders[(orders["order_status"] == "Delivered") & orders["delivery_days"].notna()]

    put(ws, 1, 1, "E-Commerce Sales & Customer Analytics: KPI summary", font=TITLE)
    put(ws, 2, 1, "Synthetic data created for portfolio use. Amounts in INR. Sales exclude Cancelled and Returned orders. "
                  "Blue = value copied from Python, black = formula.", font=NOTE)
    write_header(ws, 4, ["KPI", "Value", "How it is calculated"])
    kpis = [
        ("Total Revenue", f"=SUM(Dashboard_Data!$H$2:$H${dd_last})", MONEY, "Sum of quantity x price x (1 - discount)"),
        ("Total Profit", f"=SUM(Dashboard_Data!$J$2:$J${dd_last})", MONEY, "Revenue minus product cost"),
        ("Profit Margin", "=B6/B5", PERCENT, "Total profit / total revenue"),
        ("Total Orders", f"=Monthly_Sales!B{refs['monthly_total']}", MONEY, "Distinct valid orders (from Monthly_Sales)"),
        ("Average Order Value", "=B5/B8", MONEY, "Total revenue / total orders"),
        ("Customers with a Sale", f'=COUNTIFS(Customer_Analysis!$E${cust_first}:$E${cust_last},">=1")', MONEY, "Customers with at least 1 valid order"),
        ("Registered Customers", f"=COUNTA(Customer_Analysis!$A${cust_first}:$A${cust_last})", MONEY, "All customers in the customer table"),
        ("Repeat Customer Rate", "=Customer_Analysis!B22", PERCENT, "Customers with 2+ orders / customers with a sale"),
        ("Revenue per Customer", "=B5/B10", MONEY, "Total revenue / customers with a sale"),
        ("Average Delivery Days", round(float(delivered["delivery_days"].mean()), 2), "0.00", "Delivered orders with a valid delivery date (from Python)"),
        (f"Delivered within {DELIVERY_SLA_DAYS} Days", float((delivered["delivery_days"] <= DELIVERY_SLA_DAYS).mean()), PERCENT, "Share of those orders (from Python)"),
        ("Average Rating", round(float(reviews["rating"].mean()), 2), "0.00", "Mean of valid review ratings (from Python)"),
    ]
    for i, (label, value, number_format, how) in enumerate(kpis):
        r = 5 + i
        put(ws, r, 1, label, font=BOLD)
        put(ws, r, 2, value, number_format)
        put(ws, r, 3, how, font=NOTE)

    status_header = 19
    put(ws, status_header - 1, 1, "Order status (all orders)", font=BOLD)
    write_header(ws, status_header, ["Status", "Orders (from Python)", "% of Orders"])
    counts = orders["order_status"].value_counts()
    for i, (status, count) in enumerate(counts.items()):
        r = status_header + 1 + i
        put(ws, r, 1, status, font=BLACK)
        put(ws, r, 2, int(count), MONEY)
        put(ws, r, 3, f"=B{r}/SUM($B${status_header + 1}:$B${status_header + len(counts)})", PERCENT)
    cancel_row = status_header + len(counts) + 2
    put(ws, cancel_row, 1, "Cancellation + return rate", font=BOLD)
    put(ws, cancel_row, 3, f'=SUMIFS(C{status_header + 1}:C{status_header + len(counts)},A{status_header + 1}:A{status_header + len(counts)},"Cancelled")'
                           f'+SUMIFS(C{status_header + 1}:C{status_header + len(counts)},A{status_header + 1}:A{status_header + len(counts)},"Returned")', PERCENT)

    check_row = cancel_row + 3
    put(ws, check_row - 1, 1, "Reconciliation checks (each sheet should add up to the same total revenue)", font=BOLD)
    write_header(ws, check_row, ["Check", "Sheet total", "Result"])
    checks = [
        ("Monthly_Sales", f"=Monthly_Sales!C{refs['monthly_total']}"),
        ("Category_Analysis", f"=Category_Analysis!B{refs['category_total']}"),
        ("Geography", f"=Geography!C{refs['geo_total']}"),
        ("Customer_Analysis", "=Customer_Analysis!C10"),
        ("Product_Analysis", f"=SUM(Product_Analysis!$G${refs['prod_first']}:$G${refs['prod_last']})"),
    ]
    for i, (label, formula) in enumerate(checks):
        r = check_row + 1 + i
        put(ws, r, 1, label, font=BLACK)
        put(ws, r, 2, formula, MONEY)
        put(ws, r, 3, f'=IF(ABS(B{r}-$B$5)<1,"OK","CHECK")')
    set_widths(ws, [32, 20, 58])


def main():
    REPORTS_DIR.mkdir(exist_ok=True)
    _, _, orders, _, reviews = load_processed()
    sales = build_sales_table(valid_only=True)
    product_summary = pd.read_csv(PROCESSED_DIR / "product_summary.csv")
    customer_summary = pd.read_csv(PROCESSED_DIR / "customer_summary.csv")

    wb = Workbook()
    wb.active.title = "KPI_Summary"
    dd_last = build_dashboard_data(wb, sales)
    monthly_total = build_monthly_sales(wb, sales, dd_last)
    _, _, category_total = build_category_analysis(wb, sales, dd_last)
    prod_first, prod_last = build_product_analysis(wb, product_summary)
    cust_first, cust_last = build_customer_analysis(wb, customer_summary)
    _, _, geo_total = build_geography(wb, sales, dd_last)
    build_kpi_summary(wb, orders, reviews, {
        "dd_last": dd_last, "monthly_total": monthly_total, "category_total": category_total,
        "prod_first": prod_first, "prod_last": prod_last, "cust_first": cust_first, "cust_last": cust_last,
        "geo_total": geo_total})

    wb._sheets = [wb[name] for name in ["KPI_Summary", "Dashboard_Data", "Monthly_Sales", "Category_Analysis",
                                        "Product_Analysis", "Customer_Analysis", "Geography"]]
    wb.save(REPORTS_DIR / "ecommerce_analysis.xlsx")
    print("Saved reports/ecommerce_analysis.xlsx")


if __name__ == "__main__":
    main()
